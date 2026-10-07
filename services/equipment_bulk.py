import csv
import io
import json
from collections import defaultdict
from typing import Any
from uuid import UUID
from zipfile import BadZipFile

from fastapi import HTTPException, UploadFile, status
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from crud import equipment as crud_equipment
from crud import equipment_taxonomy as crud_taxonomy
from models.equipment import Equipment
from models.user import User
from schemas.equipment import (
    EQUIPMENT_BULK_COLUMNS,
    EquipmentBulkError,
    EquipmentBulkImportResponse,
    EquipmentCreate,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 1000


def _http_error(
    status_code: int,
    message: str,
    total_records: int = 0,
    errors: list[EquipmentBulkError] | None = None,
) -> HTTPException:
    row_errors = errors or []
    return HTTPException(
        status_code=status_code,
        detail={
            "message": message,
            "total_records": total_records,
            "successful_count": 0,
            "failed_count": len(row_errors),
            "created_ids": [],
            "errors": [error.model_dump(mode="json") for error in row_errors],
        },
    )


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def _rows_from_csv(contents: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        text = contents.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El CSV debe estar codificado en UTF-8.",
        ) from exc

    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        headers = next(reader)
    except StopIteration:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo CSV está vacío.",
        )
    _validate_headers(headers)
    rows = []
    for row_number, cells in enumerate(reader, start=2):
        if not any(_clean_cell(cell) is not None for cell in cells):
            continue
        if len(cells) != len(headers):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"La fila {row_number} no tiene la misma cantidad de columnas que los encabezados.",
            )
        if len(rows) >= MAX_IMPORT_ROWS:
            raise _http_error(
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                f"El archivo supera el máximo de {MAX_IMPORT_ROWS} equipos.",
                total_records=len(rows) + 1,
            )
        rows.append((row_number, dict(zip(headers, cells))))
    return rows


def _rows_from_xlsx(contents: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(
            io.BytesIO(contents),
            read_only=True,
            data_only=True,
        )
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo no es un Excel .xlsx válido.",
        ) from exc

    try:
        if len(workbook.worksheets) != 1:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El Excel debe contener exactamente una hoja.",
            )

        worksheet = workbook.worksheets[0]
        rows_iterator = worksheet.iter_rows(values_only=True)
        headers = next(rows_iterator, None)
        if headers is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="La hoja de Excel está vacía.",
            )
        header_values = [str(value).strip() if value is not None else "" for value in headers]
        _validate_headers(header_values)

        rows = []
        for row_number, cells in enumerate(rows_iterator, start=2):
            if not any(_clean_cell(cell) is not None for cell in cells):
                continue
            if len(cells) != len(header_values):
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"La fila {row_number} no tiene la misma cantidad de columnas que los encabezados.",
                )
            if len(rows) >= MAX_IMPORT_ROWS:
                raise _http_error(
                    status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    f"El archivo supera el máximo de {MAX_IMPORT_ROWS} equipos.",
                    total_records=len(rows) + 1,
                )
            rows.append((row_number, dict(zip(header_values, cells))))
        return rows
    finally:
        workbook.close()


def _validate_headers(headers: list[str]) -> None:
    expected = list(EQUIPMENT_BULK_COLUMNS)
    if headers != expected:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "message": "Los encabezados no coinciden con la plantilla de equipos.",
                "expected_headers": expected,
                "received_headers": headers,
            },
        )


def _parse_rows(
    rows: list[tuple[int, dict[str, Any]]],
    selected_taxonomy_id: UUID,
) -> tuple[list[dict[str, Any]], list[EquipmentBulkError]]:
    valid_records: list[dict[str, Any]] = []
    errors: list[EquipmentBulkError] = []

    if len(rows) > MAX_IMPORT_ROWS:
        raise _http_error(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            f"El archivo supera el máximo de {MAX_IMPORT_ROWS} equipos.",
            total_records=len(rows),
        )
    if not rows:
        raise _http_error(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "El archivo no contiene filas de equipos para importar.",
        )

    for row_number, raw_record in rows:
        record = {key: _clean_cell(value) for key, value in raw_record.items()}
        tag = record.get("tag_number")

        raw_taxonomy_id = record.get("taxonomy_id")
        if raw_taxonomy_id is not None:
            try:
                row_taxonomy_id = UUID(str(raw_taxonomy_id))
            except (ValueError, TypeError, AttributeError):
                row_taxonomy_id = None
            if row_taxonomy_id != selected_taxonomy_id:
                errors.append(
                    EquipmentBulkError(
                        row=row_number,
                        tag_number=str(tag) if tag is not None else None,
                        error="taxonomy_id debe estar vacío o coincidir con la taxonomía seleccionada por nombre.",
                    )
                )
                continue

        record["taxonomy_id"] = selected_taxonomy_id
        raw_specifications = record.get("technical_specifications")
        if isinstance(raw_specifications, str):
            try:
                raw_specifications = json.loads(raw_specifications)
            except json.JSONDecodeError:
                errors.append(
                    EquipmentBulkError(
                        row=row_number,
                        tag_number=str(tag) if tag is not None else None,
                        error="technical_specifications debe ser un objeto JSON válido.",
                    )
                )
                continue
            if raw_specifications is not None and not isinstance(raw_specifications, dict):
                errors.append(
                    EquipmentBulkError(
                        row=row_number,
                        tag_number=str(tag) if tag is not None else None,
                        error="technical_specifications debe ser un objeto JSON.",
                    )
                )
                continue

        record["technical_specifications"] = raw_specifications
        try:
            equipment = EquipmentCreate.model_validate(record)
        except ValidationError as exc:
            messages = "; ".join(
                f"{'.'.join(str(part) for part in issue['loc'])}: {issue['msg']}"
                for issue in exc.errors()
            )
            errors.append(
                EquipmentBulkError(
                    row=row_number,
                    tag_number=str(tag) if tag is not None else None,
                    error=messages,
                )
            )
            continue
        valid_records.append(equipment.model_dump())

    return valid_records, errors


async def import_equipments(
    db: AsyncSession,
    taxonomy_id: UUID | None,
    taxonomy_name: str | None,
    file: UploadFile,
    user: User,
) -> EquipmentBulkImportResponse:
    normalized_taxonomy_name = taxonomy_name.strip() if taxonomy_name else None
    if taxonomy_id is None and not normalized_taxonomy_name:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Debe proporcionar taxonomy_id o taxonomy_name para seleccionar la taxonomía.",
        )

    if taxonomy_id is not None:
        taxonomy = await crud_taxonomy.get_taxonomy_by_id(db, taxonomy_id)
        if taxonomy is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"No existe una taxonomía activa con el UUID '{taxonomy_id}'.",
            )
        if normalized_taxonomy_name and (
            taxonomy.name.strip().casefold() != normalized_taxonomy_name.casefold()
        ):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="taxonomy_name no corresponde al taxonomy_id indicado.",
            )
    else:
        taxonomies = await crud_taxonomy.get_taxonomies_by_name(
            db, normalized_taxonomy_name
        )
        if not taxonomies:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"No existe una taxonomía activa con el nombre '{normalized_taxonomy_name}'.",
            )
        if len(taxonomies) > 1:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    f"El nombre '{normalized_taxonomy_name}' corresponde a varias "
                    "taxonomías. Envía taxonomy_id para seleccionar una."
                ),
            )
        taxonomy = taxonomies[0]
        taxonomy_id = taxonomy.id

    contents = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"El archivo supera el límite de {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.",
        )

    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        rows = _rows_from_csv(contents)
    elif filename.endswith(".xlsx"):
        rows = _rows_from_xlsx(contents)
    else:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Formato no admitido. Sube un archivo .xlsx o .csv.",
        )

    valid_records, errors = _parse_rows(rows, taxonomy_id)
    record_tags = [record["tag_number"] for record in valid_records]
    rows_by_tag: dict[str, list[int]] = defaultdict(list)
    for row_number, raw_record in rows:
        raw_tag = _clean_cell(raw_record.get("tag_number"))
        if isinstance(raw_tag, str):
            rows_by_tag[raw_tag].append(row_number)

    for tag_number, duplicate_rows in rows_by_tag.items():
        if len(duplicate_rows) > 1:
            for row_number in duplicate_rows:
                errors.append(
                    EquipmentBulkError(
                        row=row_number,
                        tag_number=tag_number,
                        error="El tag está duplicado dentro del archivo.",
                    )
                )

    existing_tags = await crud_equipment.get_existing_equipment_tags(db, record_tags)
    for row_number, raw_record in rows:
        tag_number = _clean_cell(raw_record.get("tag_number"))
        if tag_number in existing_tags:
            errors.append(
                EquipmentBulkError(
                    row=row_number,
                    tag_number=str(tag_number),
                    error="El tag ya existe en la base de datos, incluso si está dado de baja.",
                )
            )

    if errors:
        raise _http_error(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "No se importó ningún equipo. Corrige los errores y vuelve a enviar el archivo.",
            total_records=len(rows),
            errors=errors,
        )

    records = []
    for record in valid_records:
        record["taxonomy_id"] = taxonomy_id
        record["created_by"] = user.id
        records.append(record)

    try:
        created = await crud_equipment.create_equipments_bulk(db, records)
    except IntegrityError as exc:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="No se importó ningún equipo; un tag ya existe o hubo un conflicto de integridad.",
        ) from exc

    return EquipmentBulkImportResponse(
        total_records=len(rows),
        successful_count=len(created),
        failed_count=0,
        created_ids=[equipment.id for equipment in created],
        errors=[],
    )


def _equipment_export_row(equipment: Equipment) -> list[Any]:
    specifications = equipment.technical_specifications
    return [
        str(equipment.taxonomy_id),
        equipment.tag_number,
        equipment.name,
        equipment.equipment_type,
        equipment.operational_status,
        equipment.brand,
        equipment.model,
        equipment.function_description,
        json.dumps(specifications, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
        if specifications is not None
        else None,
    ]


async def export_equipments(
    db: AsyncSession,
    taxonomy_id: UUID,
    file_format: str,
) -> tuple[bytes, str, str]:
    taxonomy = await crud_taxonomy.get_taxonomy_by_id(db, taxonomy_id)
    if taxonomy is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La taxonomía seleccionada no existe o está dada de baja.",
        )
    equipments = await crud_equipment.get_all_equipments_by_taxonomy(db, taxonomy_id)

    if file_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(EQUIPMENT_BULK_COLUMNS)
        for equipment in equipments:
            writer.writerow(_equipment_export_row(equipment))
        return (
            output.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8",
            f"equipos_{taxonomy_id}.csv",
        )

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Equipos")
    worksheet.append(list(EQUIPMENT_BULK_COLUMNS))
    for equipment in equipments:
        worksheet.append(_equipment_export_row(equipment))
    output = io.BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"equipos_{taxonomy_id}.xlsx",
    )
